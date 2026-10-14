#!/usr/bin/env python3
"""
Generates Methodology_Revised.docx from the updated Chapter 3 specifications.
Applies professional academic styling, standard typography, and clean tables.
"""

from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    v_border = OxmlElement('w:insideV')
    v_border.set(qn('w:val'), 'none')
    tblBorders.append(v_border)
    tblPr.append(tblBorders)

def create_methodology_doc(output_path: Path):
    doc = docx.Document()

    # Page Margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(11.5)
    style_normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)

    # Document Header
    p_meta = doc.add_paragraph()
    r_meta = p_meta.add_run("AGRICARE: DIGITAL SOIL MAPPING & PRECISION AGRONOMIC FRAMEWORK")
    r_meta.font.name = 'Arial'
    r_meta.font.size = Pt(9.5)
    r_meta.font.bold = True
    r_meta.font.color.rgb = RGBColor(0x1B, 0x4F, 0x72)
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # Chapter Title
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("CHAPTER 3: PROPOSED METHODOLOGY")
    r1.font.name = 'Arial'
    r1.font.size = Pt(20)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0x15, 0x43, 0x60)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(14)

    # Section 3.1
    h2 = doc.add_heading(level=2)
    r2 = h2.add_run("3.1 Detailed Explanation of System Architecture Schematic")
    r2.font.name = 'Arial'
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(0x1A, 0x52, 0x76)

    # 3.1.1 Study Area
    h3 = doc.add_heading(level=3)
    r3 = h3.add_run("3.1.1 Study Area and Open Geospatial Data Ingestion")
    r3.font.name = 'Arial'
    r3.font.size = Pt(12)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(0x21, 0x61, 0x8C)

    doc.add_paragraph(
        "The study area comprises SPSR Nellore District, Andhra Pradesh, a coastal agro-ecological zone in peninsular India "
        "characterized by contrasting physiography: an irrigated deltaic alluvial plain dominated by double-cropped wetland paddy, "
        "flanked by rainfed semi-arid uplands cultivated with cotton, groundnut, and maize. Ground-truth soil chemistry originates from "
        "909 geo-referenced Soil Health Card (SHC) laboratory wet-chemistry records spanning 46 mandals. Four target soil properties "
        "are predicted: Available Nitrogen (N), Available Phosphorus (P), and Available Potassium (K), all measured in kg/ha, alongside "
        "Soil Organic Carbon (OC, %). For cross-season transferability benchmarks, three multi-year Harmonized Sentinel-2 SCORPAN datasets "
        "are utilized for the 2023–24 (948 samples), 2024–25 (952 samples), and 2025–26 (938 samples) agricultural cycles."
    )
    doc.add_paragraph(
        "Explanatory environmental covariates are derived from four open-access geospatial repositories on Google Earth Engine (GEE): "
        "(1) Sentinel-2 MSI Level-2A surface reflectance (12 bands); (2) NASA SRTM 30m Digital Elevation Model; "
        "(3) ECMWF ERA5-Land reanalysis and TerraClimate (7 hydrometeorological variables); and (4) ISRIC SoilGrids 250m mineral fractions "
        "(clay, silt, and sand in g/kg). Cloud-native processing in GEE enables instantaneous covariate extraction across the district without local data storage."
    )

    # 3.1.2 5-Stage Pipeline
    h3_arch = doc.add_heading(level=3)
    r_arch = h3_arch.add_run("3.1.2 Five-Stage Operational Pipeline Architecture")
    r_arch.font.name = 'Arial'
    r_arch.font.size = Pt(12)
    r_arch.font.bold = True
    r_arch.font.color.rgb = RGBColor(0x21, 0x61, 0x8C)

    doc.add_paragraph(
        "AgriCare is structured as a sequential 5-stage lifecycle, operationalizing Hans Jenny's pedometric state factor equation "
        "as formalized by McBratney's SCORPAN paradigm: S = f(s, c, o, r, p, a, n). The five pipeline stages are:"
    )

    stages = [
        ("Stage 1 (Ground-Truth Digitization): ", "Paper Soil Health Cards are parsed using OpenCV morphological grid filtering, adaptive Gaussian thresholding, and bilingual EasyOCR, verified by regex and physical boundary filters (N=909 records)."),
        ("Stage 2 (GEE Bare-Soil Covariate Extraction): ", "Sentinel-2 L2A surface reflectance is filtered during the pre-sowing window (April 1 to June 30) using BSI > 0.05 and NDVI < 0.25 bare-soil masks, resampled to a uniform 10m grid, and sampled alongside SRTM DEM, ERA5-Land, and SoilGrids layers."),
        ("Stage 3 (SCORPAN Feature Engineering): ", "Synthesizes 39 baseline covariates and 8 interaction/trend terms into a 47-dimensional feature matrix, incorporating sediment redistribution and catenary hydrology (Clay_x_Elevation, Moisture_div_Slope)."),
        ("Stage 4 (Dual-Head Validation, Few-Shot Domain Calibration & CQR): ", "Validates generalization via Random 5-Fold, Spatial-Block K-Fold (5km grids), and Out-of-Time testing. Cross-season monsoon drift is neutralized via N=25 anchor calibration, while Conformalized Quantile Regression (CQR) guarantees >=90% empirical coverage."),
        ("Stage 5 (ANGRAU Precision Fertilizer Advisory): ", "Integrates continuous hybrid VotingRegressor forecasts with ICAR ExtraTreesClassifier ratings into ANGRAU STCR mass-balance equations, automating plot-specific Urea, DAP, and MOP recommendations (kg/acre).")
    ]
    for st_title, st_desc in stages:
        p_st = doc.add_paragraph(style='List Bullet')
        r_t = p_st.add_run(st_title)
        r_t.bold = True
        p_st.add_run(st_desc)

    # Section 3.2
    h2_mod = doc.add_heading(level=2)
    r_mod = h2_mod.add_run("3.2 Modular Breakdown of the Technical Pipeline")
    r_mod.font.name = 'Arial'
    r_mod.font.size = Pt(14)
    r_mod.font.bold = True
    r_mod.font.color.rgb = RGBColor(0x1A, 0x52, 0x76)

    # Table 3.1
    p_t1 = doc.add_paragraph()
    r_t1 = p_t1.add_run("Table 3.1: Spectral Band Specifications of the Multispectral Instrument (Effective Resolution: 10m)")
    r_t1.bold = True
    r_t1.font.size = Pt(11)

    table1_data = [
        ("Band Identifier", "Central λ (nm)", "Bandwidth (nm)", "Spatial Resolution", "Physical Pedological Sensitivity"),
        ("B1 Coastal Aerosol", "443", "20", "10 m (Resampled from 60m)", "Atmospheric aerosol scattering correction"),
        ("B2 Blue", "490", "65", "10 m (Native)", "Organic matter absorption; iron oxide discrimination"),
        ("B3 Green", "560", "35", "10 m (Native)", "Chlorophyll reflectance peak; green peak"),
        ("B4 Red", "665", "30", "10 m (Native)", "Chlorophyll and ferric iron (hematite) absorption"),
        ("B5 Red Edge 1", "705", "15", "10 m (Resampled from 20m)", "Boundary of vegetation stress & canopy chlorophyll"),
        ("B6 Red Edge 2", "740", "15", "10 m (Resampled from 20m)", "Leaf cellular structure; surrogate for nitrogen"),
        ("B7 Red Edge 3", "783", "20", "10 m (Resampled from 20m)", "Canopy biomass density & Leaf Area Index"),
        ("B8 Broad NIR", "842", "115", "10 m (Native)", "Structural scattering by vegetation canopy"),
        ("B8A Narrow NIR", "865", "20", "10 m (Resampled from 20m)", "Narrow near-infrared plateau; organic carbon"),
        ("B9 Water Vapour", "945", "20", "10 m (Resampled from 60m)", "Atmospheric water-vapour column correction"),
        ("B11 SWIR 1", "1610", "90", "10 m (Resampled from 20m)", "Clay-mineral (O-H) lattice absorption & moisture"),
        ("B12 SWIR 2", "2190", "180", "10 m (Resampled from 20m)", "C-H absorption in organic carbon; carbonates")
    ]
    t1 = doc.add_table(rows=len(table1_data), cols=5)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1)
    for r_idx, row in enumerate(table1_data):
        for c_idx, val in enumerate(row):
            cell = t1.cell(r_idx, c_idx)
            cell.text = val
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx != 1 and c_idx != 2 else WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.name = 'Times New Roman'
            if r_idx == 0:
                set_cell_background(cell, "1B4F72")
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif r_idx % 2 == 1:
                set_cell_background(cell, "F2F4F4")

    doc.add_paragraph(
        "Note on Spatial Resampling: In Google Earth Engine, all 20m and 60m Sentinel-2 Level-2A surface reflectance bands "
        "are resampled to a uniform 10m processing resolution via bilinear interpolation to align with native 10m visible bands "
        "(B2, B3, B4, B8), enabling fine-scale 10m plot-level digital soil mapping."
    )

    # Table 3.2
    p_t2 = doc.add_paragraph()
    r_t2 = p_t2.add_run("Table 3.2: Complete Catalog of 47 Engineered SCORPAN Features in AgriCare")
    r_t2.bold = True
    r_t2.font.size = Pt(11)

    table2_data = [
        ("Category", "Count", "Variable Names / Symbols", "Physical Domain Rationale"),
        ("1. Sentinel-2 Bands", "12", "B1, B2, B3, B4, B5, B6, B7, B8, B8A, B9, B11, B12", "Direct optical and shortwave infrared reflectance (effective 10m)"),
        ("2. Spectral Indices", "8", "NDVI, NDRE, EVI, SAVI, BSI, Clay_Ratio, OC_Index, Integral_NDVI", "Vegetation vigor, bare soil albedo, clay/carbon absorption"),
        ("3. Geomorphometry", "7", "Elevation_m, Slope_deg, Hillshade, TWI, Curvature, Aspect_Sin, Aspect_Cos", "Topographic water accumulation, illumination, drainage"),
        ("4. Climate Reanalysis", "7", "Soil_Moisture, Soil_Temp, LST_K, Precip_Mean, PET, Aridity_Index, VPD", "Thermodynamic soil moisture, weathering, evapotranspiration"),
        ("5. Texture & Space", "5", "Clay_Fraction, Sand_Fraction, Silt_Fraction, Latitude, Longitude", "Physical mineral matrix, cation exchange, coordinates"),
        ("6. Physical Interactions", "3", "Clay_x_Moisture, Temp_x_VPD, BSI_div_NDVI", "Coupled hydrological and atmospheric evaporative demand"),
        ("7. Spatial Polynomials", "3", "Spatial_Lat2, Spatial_Lon2, Spatial_Lat_Lon", "Geographic trend surface across Nellore delta"),
        ("8. Soil-Terrain Hydrology", "2", "Clay_x_Elevation, Moisture_div_Slope", "Sediment redistribution & erosion-deposition dynamics")
    ]
    t2 = doc.add_table(rows=len(table2_data), cols=4)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t2)
    for r_idx, row in enumerate(table2_data):
        for c_idx, val in enumerate(row):
            cell = t2.cell(r_idx, c_idx)
            cell.text = val
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx != 1 else WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.name = 'Times New Roman'
            if r_idx == 0:
                set_cell_background(cell, "1B4F72")
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif r_idx % 2 == 1:
                set_cell_background(cell, "F2F4F4")

    doc.add_paragraph(
        "Ablation Impact: Incremental ablation experiments prove that Sentinel-2 optical bands alone achieve a mean R² of 0.3415. "
        "Layering geomorphometry (R²=0.4120), climate reanalysis (R²=0.4682), soil texture (R²=0.4935), and the 8 physical interaction terms "
        "(Categories 6, 7, and 8) elevates overall multi-nutrient mean R² to 0.5207."
    )

    # Head A & Head B Architecture
    h3_head = doc.add_heading(level=3)
    r_head = h3_head.add_run("3.2.4 Dual-Head Machine Learning Architecture")
    r_head.font.name = 'Arial'
    r_head.font.size = Pt(12)
    r_head.font.bold = True
    r_head.font.color.rgb = RGBColor(0x21, 0x61, 0x8C)

    doc.add_paragraph(
        "To eliminate boundary thresholding errors inherent in monolithic regression models, AgriCare introduces a parallel Dual-Head architecture:\n"
        "• Head A (Continuous Regression): Formulated as a weighted hybrid VotingRegressor combining 70% ExtraTreesRegressor (350 trees, max_depth=16, min_samples_leaf=2) "
        "and 30% GradientBoostingRegressor:\n"
        "    y_hat_continuous = 0.70 * y_hat_ExtraTrees + 0.30 * y_hat_GradientBoosting    ... (3.19)\n"
        "Target variables P and K undergo logarithmic transformation prior to fitting to eliminate right-skewness:\n"
        "    y_trans = ln(1 + y)  ==>  y_hat = max( 0, exp(y_hat_trans) - 1 )    ... (3.18)\n"
        "• Head B (ICAR Fertility Classification): An independent ExtraTreesClassifier (300 trees, max_depth=14) trained directly on statutory "
        "Indian Council of Agricultural Research (ICAR) fertility classes (Low, Medium, High). Predictions output posterior confidence percentages:\n"
        "    Confidence(%) = max_c P(c | x) * 100    ... (3.22)"
    )

    # Table 3.3
    p_t3 = doc.add_paragraph()
    r_t3 = p_t3.add_run("Table 3.3: Statutory ICAR Soil Fertility Classification Rating Scales")
    r_t3.bold = True
    r_t3.font.size = Pt(11)

    table3_data = [
        ("Soil Nutrient / Property", "Low (Tier 0)", "Medium (Tier 1)", "High (Tier 2)", "Analytical Method"),
        ("Available Nitrogen (N)", "< 280 kg/ha", "280 – 560 kg/ha", "> 560 kg/ha", "Alkaline KMnO4 Distillation"),
        ("Available Phosphorus (P, Olsen)", "< 10 kg/ha (equiv. < 23 kg/ha P2O5)", "10 – 25 kg/ha (equiv. 23 – 56 kg/ha P2O5)", "> 25 kg/ha (equiv. > 56 kg/ha P2O5)", "Olsen 0.5M NaHCO3 Spectrophotometry"),
        ("Available Potassium (K)", "< 140 kg/ha (equiv. < 168 kg/ha K2O)", "140 – 280 kg/ha (equiv. 168 – 337 kg/ha K2O)", "> 280 kg/ha (equiv. > 337 kg/ha K2O)", "Neutral 1N NH4OAc Flame Photometry"),
        ("Soil Organic Carbon (OC)", "< 0.50 %", "0.50 % – 0.75 %", "> 0.75 %", "Walkley-Black Wet Oxidation")
    ]
    t3 = doc.add_table(rows=len(table3_data), cols=5)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t3)
    for r_idx, row in enumerate(table3_data):
        for c_idx, val in enumerate(row):
            cell = t3.cell(r_idx, c_idx)
            cell.text = val
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx != 1 and c_idx != 2 and c_idx != 3 else WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.size = Pt(9.5)
            p.runs[0].font.name = 'Times New Roman'
            if r_idx == 0:
                set_cell_background(cell, "1B4F72")
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif r_idx % 2 == 1:
                set_cell_background(cell, "F2F4F4")

    # Few-Shot & CQR
    h3_val = doc.add_heading(level=3)
    r_val = h3_val.add_run("3.2.5 Few-Shot Domain Adaptation & Conformalized Quantile Regression")
    r_val.font.name = 'Arial'
    r_val.font.size = Pt(12)
    r_val.font.bold = True
    r_val.font.color.rgb = RGBColor(0x21, 0x61, 0x8C)

    doc.add_paragraph(
        "Few-Shot Seasonal Calibration: Inter-annual monsoon variance alters baseline soil albedo and moisture, introducing cross-season "
        "bias (MBE = +19.74 kg/ha for N, +80.82 kg/ha for K). AgriCare draws a minimal batch of K=25 anchor ground-truth tests from the new season, "
        "calculates the mean residual offset Delta = (1/K) * Sum(y_true - y_hat_base), and shifts un-sampled predictions:\n"
        "    y_hat_calibrated = max(0, y_hat_base + Delta)    ... (3.27)\n"
        "This additive offset eliminates retraining overhead, reduces cross-cycle RMSE by 32.4%, and recovers Nitrogen R² from 0.3420 back to 0.5840."
    )
    doc.add_paragraph(
        "Conformalized Quantile Regression (CQR): Raw Quantile Random Forest intervals under-cover test observations (PICP ≈ 73%–77% on uncalibrated sets). "
        "CQR (Romano et al., NeurIPS 2019) evaluates non-conformity scores on a 20% holdout calibration set (N_calib ≈ 180):\n"
        "    E_i = max( q_lo(x_i) - y_i,  y_i - q_hi(x_i) )    ... (3.28)\n"
        "The empirical correction quantile q_hat_conf is evaluated at level ceil((N_calib + 1)(1 - alpha)) / N_calib. Expanding test intervals by q_hat_conf:\n"
        "    Interval_calibrated = [ max(0, q_lo - q_hat_conf),  q_hi + q_hat_conf ]    ... (3.29)\n"
        "guarantees empirical prediction coverage PICP >= 90.0% (achieving 95.10% for P, 92.86% for OC, and >=90% under calibrated domain drift)."
    )

    # ANGRAU Advisory Engine
    h3_adv = doc.add_heading(level=3)
    r_adv = h3_adv.add_run("3.2.6 ANGRAU Precision Fertilizer Advisory Engine")
    r_adv.font.name = 'Arial'
    r_adv.font.size = Pt(12)
    r_adv.font.bold = True
    r_adv.font.color.rgb = RGBColor(0x21, 0x61, 0x8C)

    doc.add_paragraph(
        "Commercial fertilizer prescriptions are automated using Acharya N.G. Ranga Agricultural University (ANGRAU) STCR baseline recommendations "
        "for major crops in SPSR Nellore (Paddy: 120-60-40, Maize: 150-60-50, Cotton: 120-60-60, Groundnut: 20-40-50 in N-P2O5-K2O kg/ha). "
        "Requirements are adjusted by statutory tier multipliers (+25% for Low, 0% for Medium, -25% for High). "
        "Because Di-Ammonium Phosphate (DAP) supplies 18% N and 46% P2O5, DAP is calculated first, and its nitrogen contribution is credited against Urea:\n"
        "    DAP (kg/ha) = R_P / 0.46;    N_DAP = 0.18 * DAP    ... (3.31)\n"
        "    Urea (kg/ha) = max(0, R_N - N_DAP) / 0.46;    MOP (kg/ha) = R_K / 0.60    ... (3.32)\n"
        "    Prescription (kg/acre) = Dose (kg/ha) * 0.4047    ... (3.33)\n"
        "For Paddy under Low-N, Medium-P, High-K conditions, the exact prescription is 52.8 kg DAP, 111.3 kg Urea, and 20.2 kg MOP per acre."
    )

    # Multi-Sensor & Micronutrients
    h3_ext = doc.add_heading(level=3)
    r_ext = h3_ext.add_run("3.2.7 Future Multi-Sensor Radar & Micronutrient Extensions")
    r_ext.font.name = 'Arial'
    r_ext.font.size = Pt(12)
    r_ext.font.bold = True
    r_ext.font.color.rgb = RGBColor(0x21, 0x61, 0x8C)

    doc.add_paragraph(
        "To resolve optical penetration barriers for non-reflective phosphate ions and secondary nutrients, three extensions are detailed:\n"
        "1. Sentinel-1 C-Band SAR Radar Fusion: Ingestion of dual-polarization backscatter (VV, VH in dB) and radar indices (Ratio = VH/VV, "
        "RVI = 4*VH / (VV + VH)) to penetrate vegetative canopy and directly gauge topsoil dielectric moisture.\n"
        "2. Spaceborne Hyperspectral Ingestion (PRISMA / EnMAP): Contiguous 240-band spectroscopy (400–2500 nm) to detect narrow mineral absorption doublets.\n"
        "3. 12-Parameter Soil Health Card Expansion: Multi-label classification heads for secondary (S, Ca, Mg) and micronutrients (Zn, Fe, B), "
        "eliminating crop yield bottlenecks governed by Liebig's Law of the Minimum."
    )

    doc.save(str(output_path))
    print(f"[+] Successfully generated Word document at: {output_path}")

if __name__ == "__main__":
    out_file1 = Path("c:/Users/Vidya sagar/Sunny/projects/Temp work/AgriCare/Methodology_Revised.docx")
    out_file2 = Path("c:/Users/Vidya sagar/Sunny/projects/Temp work/Methodology_Revised.docx")
    create_methodology_doc(out_file1)
    create_methodology_doc(out_file2)
